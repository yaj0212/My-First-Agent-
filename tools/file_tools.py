import csv
import os
import tempfile
from langchain_core.tools import tool

OUTPUT_DIR = os.path.join(tempfile.gettempdir(), "my-first-agent")
os.makedirs(OUTPUT_DIR, exist_ok=True)


def _make_path(filename: str, ext: str) -> str:
    if not filename.endswith(ext):
        filename = filename + ext
    return os.path.join(OUTPUT_DIR, filename)


@tool
def generate_csv(filename: str, headers: list[str], rows: list[list]) -> str:
    """Generate a CSV file with the given headers and rows. Returns the file path."""
    if not filename:
        return "Error: filename cannot be empty."
    if not headers:
        return "Error: headers cannot be empty."
    path = _make_path(filename, ".csv")
    with open(path, "w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(rows)
    return f"FILE_SAVED:{path}"


@tool
def generate_markdown(filename: str, content: str) -> str:
    """Generate a Markdown (.md) file with the given content. Returns the file path."""
    if not filename:
        return "Error: filename cannot be empty."
    if not content:
        return "Error: content cannot be empty."
    path = _make_path(filename, ".md")
    with open(path, "w") as f:
        f.write(content)
    return f"FILE_SAVED:{path}"


@tool
def generate_python_file(filename: str, code: str) -> str:
    """Generate a Python (.py) file with the given code. Returns the file path."""
    if not filename:
        return "Error: filename cannot be empty."
    if not code:
        return "Error: code cannot be empty."
    path = _make_path(filename, ".py")
    with open(path, "w") as f:
        f.write(code)
    return f"FILE_SAVED:{path}"


@tool
def generate_pdf(filename: str, content: str) -> str:
    """Generate a PDF file with the given text content. Returns the file path."""
    if not filename:
        return "Error: filename cannot be empty."
    if not content:
        return "Error: content cannot be empty."
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph
        from reportlab.lib.styles import getSampleStyleSheet
    except ImportError:
        return "Error: reportlab is not installed. Run: pip install reportlab"

    path = _make_path(filename, ".pdf")
    doc = SimpleDocTemplate(path, pagesize=letter)
    styles = getSampleStyleSheet()
    paragraphs = [Paragraph(line or "&nbsp;", styles["Normal"]) for line in content.splitlines()]
    doc.build(paragraphs)
    return f"FILE_SAVED:{path}"


@tool
def generate_excel(filename: str, headers: list[str], rows: list[list]) -> str:
    """Generate an Excel (.xlsx) file with the given headers and rows. Returns the file path."""
    if not filename:
        return "Error: filename cannot be empty."
    if not headers:
        return "Error: headers cannot be empty."
    try:
        import openpyxl
    except ImportError:
        return "Error: openpyxl is not installed. Run: pip install openpyxl"

    path = _make_path(filename, ".xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(path)
    return f"FILE_SAVED:{path}"


@tool
def generate_word_doc(filename: str, content: str) -> str:
    """Generate a Word (.docx) file with the given text content. Returns the file path."""
    if not filename:
        return "Error: filename cannot be empty."
    if not content:
        return "Error: content cannot be empty."
    try:
        from docx import Document
    except ImportError:
        return "Error: python-docx is not installed. Run: pip install python-docx"

    path = _make_path(filename, ".docx")
    doc = Document()
    for line in content.splitlines():
        doc.add_paragraph(line)
    doc.save(path)
    return f"FILE_SAVED:{path}"


@tool
def generate_notebook(filename: str, cells: list[str]) -> str:
    """Generate a Jupyter Notebook (.ipynb) where each item in cells becomes a code cell."""
    if not filename:
        return "Error: filename cannot be empty."
    if not cells:
        return "Error: cells cannot be empty."
    try:
        import nbformat
    except ImportError:
        return "Error: nbformat is not installed. Run: pip install nbformat"

    path = _make_path(filename, ".ipynb")
    nb = nbformat.v4.new_notebook()
    nb.cells = [nbformat.v4.new_code_cell(cell) for cell in cells]
    with open(path, "w") as f:
        nbformat.write(nb, f)
    return f"FILE_SAVED:{path}"
